from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.importers.xlsx_importer import (
    create_import_hash,
    extract_rfid,
    import_xlsx,
    parse_duration,
)


HEADERS = [
    "Startdatum",
    "Status",
    "Dauer",
    "Gesamte Energie (kWh)",
    "MID-zertifiziert",
    "Solarstrom (kWh)",
    "Solares Verhältnis",
    "Authentifizierung",
    "Ladestation",
]


def create_test_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active

    sheet.append(HEADERS)

    sheet.append(
        [
            "15.07.2026 11:26:37",
            "Beendet (keine Aktion)",
            "0h 0m 52s",
            0,
            "Ja",
            0,
            "0 %",
            "Wartung Gottschalk (DE0D02A6)",
            "WB5",
        ]
    )

    sheet.append(
        [
            datetime(2026, 6, 9, 17, 8, 26),
            "Beendet",
            "0h 0m 50s",
            0.02,
            "Ja",
            0.02,
            "95,24 %",
            "Keine Authentifizierung",
            "WB2",
        ]
    )

    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize(
    ("duration", "expected"),
    [
        ("0h 0m 52s", timedelta(seconds=52)),
        ("0h 2m 1s", timedelta(minutes=2, seconds=1)),
        ("1h 23m 45s", timedelta(hours=1, minutes=23, seconds=45)),
        ("2h", timedelta(hours=2)),
        ("15m", timedelta(minutes=15)),
        ("9s", timedelta(seconds=9)),
    ],
)
def test_parse_duration(
    duration: str,
    expected: timedelta,
) -> None:
    assert parse_duration(duration) == expected


@pytest.mark.parametrize(
    "duration",
    [
        "",
        " ",
        "unbekannt",
        "1 Stunde",
        "12:34",
    ],
)
def test_parse_duration_rejects_invalid_values(
    duration: str,
) -> None:
    with pytest.raises(ValueError):
        parse_duration(duration)


@pytest.mark.parametrize(
    ("authentication", "expected"),
    [
        ("Wartung Gottschalk (DE0D02A6)", "DE0D02A6"),
        ("RFID Nr.4 (6005E5AC)", "6005E5AC"),
        ("RFID N.3 (D0BEEEAC)", "D0BEEEAC"),
        ("RFID Nr.5 (e033f68c)", "E033F68C"),
        ("Keine Authentifizierung", None),
        ("", None),
        ("Unbekannter Text", None),
    ],
)
def test_extract_rfid(
    authentication: str,
    expected: str | None,
) -> None:
    assert extract_rfid(authentication) == expected


def test_create_import_hash_is_stable() -> None:
    start_time = datetime(2026, 7, 15, 11, 26, 37)

    first_hash = create_import_hash(
        start_time=start_time,
        station="WB5",
        energy=1.234,
        rfid="DE0D02A6",
    )

    second_hash = create_import_hash(
        start_time=start_time,
        station="WB5",
        energy=1.234,
        rfid="DE0D02A6",
    )

    assert first_hash == second_hash
    assert len(first_hash) == 64


def test_create_import_hash_changes_with_session_data() -> None:
    start_time = datetime(2026, 7, 15, 11, 26, 37)

    first_hash = create_import_hash(
        start_time=start_time,
        station="WB5",
        energy=1.234,
        rfid="DE0D02A6",
    )

    second_hash = create_import_hash(
        start_time=start_time,
        station="WB4",
        energy=1.234,
        rfid="DE0D02A6",
    )

    assert first_hash != second_hash


def test_import_xlsx_parses_sessions(
    tmp_path: Path,
) -> None:
    xlsx_path = tmp_path / "hager-export.xlsx"
    create_test_workbook(xlsx_path)

    sessions = import_xlsx(xlsx_path)

    assert len(sessions) == 2

    first_session = sessions[0]

    assert first_session["start_time"] == datetime(
        2026,
        7,
        15,
        11,
        26,
        37,
    )
    assert first_session["end_time"] == datetime(
        2026,
        7,
        15,
        11,
        27,
        29,
    )
    assert first_session["station_id"] == "WB5"
    assert first_session["rfid"] == "DE0D02A6"
    assert first_session["energy_total_kwh"] == 0.0
    assert first_session["energy_pv_kwh"] == 0.0
    assert first_session["source"] == "xlsx"
    assert len(first_session["import_hash"]) == 64

    second_session = sessions[1]

    assert second_session["start_time"] == datetime(
        2026,
        6,
        9,
        17,
        8,
        26,
    )
    assert second_session["end_time"] == datetime(
        2026,
        6,
        9,
        17,
        9,
        16,
    )
    assert second_session["station_id"] == "WB2"
    assert second_session["rfid"] is None
    assert second_session["energy_total_kwh"] == pytest.approx(0.02)
    assert second_session["energy_pv_kwh"] == pytest.approx(0.02)


def test_import_xlsx_rejects_missing_file(
    tmp_path: Path,
) -> None:
    missing_path = tmp_path / "missing.xlsx"

    with pytest.raises(
        FileNotFoundError,
        match="XLSX-Datei nicht gefunden",
    ):
        import_xlsx(missing_path)


def test_import_xlsx_reports_invalid_duration(
    tmp_path: Path,
) -> None:
    xlsx_path = tmp_path / "invalid-duration.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    sheet.append(
        [
            "15.07.2026 11:26:37",
            "Beendet",
            "ungültig",
            1.0,
            "Ja",
            0.0,
            "0 %",
            "RFID Nr.1 (6AA972EA)",
            "WB2",
        ]
    )
    workbook.save(xlsx_path)
    workbook.close()

    with pytest.raises(
        ValueError,
        match="Fehler in XLSX-Zeile 2",
    ):
        import_xlsx(xlsx_path)


def test_import_xlsx_rejects_invalid_headers(
    tmp_path: Path,
) -> None:
    xlsx_path = tmp_path / "invalid-headers.xlsx"

    workbook = Workbook()
    sheet = workbook.active

    sheet.append(
        [
            "Datum",
            "Status",
            "Dauer",
            "Energie",
        ]
    )

    workbook.save(xlsx_path)
    workbook.close()

    with pytest.raises(
        ValueError,
        match="Ungültige XLSX-Kopfzeile",
    ):
        import_xlsx(xlsx_path)
